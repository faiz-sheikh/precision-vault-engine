#!/usr/bin/env python3
"""
curriculum_mapper.py — Phase 0 — FINAL CONSOLIDATED BUILD

This version folds in EVERY fix that was previously applied as a separate
manual patch script directly against a hand-built manifest. Those patches
existed only outside this file — meaning a fresh run of this script would
have silently regressed to the old broken behaviour. That gap is now closed:
this file is the single source of truth for Phase 0.

CONSOLIDATED FIXES (chronological order they were discovered):
  1.  Hash-suffixed IDs — slug[:45] + md5[:8] — collision-proof even when
      two titles share a 45-char prefix.
  2.  normalize_title() — strips and/or/&/etc before hashing, so "X and Y"
      and "X & Y" collapse to the same semantic ID instead of duplicating.
  3.  Two-pass dedup: Pass 1 (exact ID) + Pass 2 (semantic via normalize_title)
      — catches the slug-divergence case Pass 1 alone misses.
  4.  Word-boundary regex on all domain corrections — naive substring 'in'
      checks previously matched 'dynamics' inside 'thermodynamics'.
  5.  12_NaturalSciences + 13_MechanicalEngineering domains (was dumping into
      00_Uncategorised).
  6.  robust_json_parse() + truncate-at-line + shrink-30%-per-retry — handles
      any truncated/malformed Gemini output without wasting a retry on an
      identical oversized prompt that will fail the same way again.
  7.  core_subtopics + suggested_chapters in schema — feeds Phase 2's
      chapter-splitting synthesis (Claude's real 8,192-token output ceiling
      needs this for any module with >4-5 dense subtopics).
  8.  Canonical `subject` mapping + `track` field — subject is now always one
      of ~13 stable categories; the original descriptive label (CSE Minor,
      MDM, Honors Track, etc.) is preserved separately in `track`.
  9.  Anchor-prefixed web_queries — Gemini's raw output is DISCARDED and
      REPLACED deterministically: ambiguous bare-word queries ("nodes",
      "transitive") are never trusted from the LLM; Python always prepends
      a disambiguating anchor (acronym or title) before the subtopic phrase.
  10. Software/systems title override for engine — courses about ROS,
      microcontrollers, interfacing, operating systems, "vision system" are
      forced to remotion even when their domain defaults to manim (these are
      command-line/diagram-heavy, not math-derivation-heavy).
  11. Ceiling suggested_chapters = max(1, ceil(len(core_subtopics)/2)) — was
      floor division, under-allocating chapters for 5-subtopic modules.
  12. prerequisite_context — final cross-entry pass (after all files parsed)
      linking each module to earlier-tier/earlier-grade entries in the SAME
      canonical subject, enabling Jr -> Sr -> B.Tech continuous-flow synthesis.

API-BLEED PREVENTION (unchanged, still critical):
  - ONE Gemini call per source file (not per module) — schema-constrained
    JSON output (response_mime_type + response_schema = constrained token
    decoding, AFC cannot trigger).
  - thinking_config(thinking_budget=0) — no wasted reasoning tokens.
  - tool_config(mode=NONE) — explicit second-layer AFC prevention.
  - ALL fixes 8-12 above are deterministic Python post-processing — ZERO
    additional API calls. Never re-prompt Gemini to "fix" formatting; enforce
    every rule in code after the single call returns.
"""
import re
import json
import time
import math
import hashlib
import argparse
from pathlib import Path
from typing import Optional

try:
    from google import genai as google_genai
    from google.genai import types as gt
except ImportError:
    import sys; sys.exit("pip install google-genai")

from config import (
    log, console, get_env, safe_slug, robust_json_parse,
    MASTER_MANIFEST, STAGING_DIR, CURRICULUM_DIR, DOMAIN_KEYWORDS,
    GEMINI_FLASH_MODEL, MAX_RETRIES, calc_backoff, safe_gemini_text,
)

# ── Valid domains ───────────────────────────────────────────────────────────────
_VALID_DOMAINS = [
    "01_Math", "02_Physics", "03_Programming", "04_Electronics",
    "05_Robotics", "06_AI_ML", "07_Control", "08_Design",
    "09_Defence", "10_HSSM", "11_Projects",
    "12_NaturalSciences", "13_MechanicalEngineering",
    "00_Uncategorised",
]
MANIM_DOMAINS = {"01_Math","02_Physics","07_Control","04_Electronics","05_Robotics"}

# ── FIX 4: domain corrections, word-boundary regex (no substring collisions) ───
_DOMAIN_CORRECTIONS: list[tuple[str, str, Optional[str]]] = [
    (r'\bbiology\b',              "12_NaturalSciences",        None),
    (r'\bbotany\b',               "12_NaturalSciences",        None),
    (r'\bzoology\b',               "12_NaturalSciences",        None),
    (r'\bbiochemistry\b',         "12_NaturalSciences",        None),
    (r'\becology\b',              "12_NaturalSciences",        None),
    (r'\bgenetics\b',             "12_NaturalSciences",        None),
    (r'\blife processes\b',       "12_NaturalSciences",        None),
    (r'\bcell structure\b',       "12_NaturalSciences",        None),
    (r'\bcell biology\b',         "12_NaturalSciences",        None),
    (r'\btissues\b',              "12_NaturalSciences",        None),
    (r'\borganic chemistry\b',    "12_NaturalSciences",        None),
    (r'\binorganic chemistry\b',  "12_NaturalSciences",        None),
    (r'\bacids,?\s*bases\b',      "12_NaturalSciences",        None),
    (r'\bchemical dynamics\b',    "12_NaturalSciences",        None),
    (r'\benviro\w+ science\b',    "12_NaturalSciences",        None),
    (r'\bmicrobiology\b',         "12_NaturalSciences",        None),
    (r'\bphysiology\b',           "12_NaturalSciences",        None),
    (r'\breproduct\w*\b',         "12_NaturalSciences",        None),
    (r'\btheory of machines\b',   "13_MechanicalEngineering",  r'\b(robot|autonomous|ros)\b'),
    (r'\bkinematics of machine',  "13_MechanicalEngineering",  r'\b(robot|manipulator)\b'),
    (r'\belasticity\b',           "13_MechanicalEngineering",  None),
    (r'\bfluid mechanics\b',      "13_MechanicalEngineering",  None),
    (r'\bheat transfer\b',        "13_MechanicalEngineering",  None),
    (r'\bmanufacturing process',  "13_MechanicalEngineering",  None),
    (r'\bstrength of material',   "13_MechanicalEngineering",  None),
    (r'\bmachine design\b',       "13_MechanicalEngineering",  r'\b(robot|neural)\b'),
    (r'\bturbomachinery\b',       "13_MechanicalEngineering",  None),
    (r'\bvehicle dynamics\b',     "07_Control",                None),
    (r'\bhuman.computer interac', "03_Programming",            r'\b(defence|weapon|military)\b'),
    (r'\bhuman.machine interac',  "03_Programming",            r'\b(defence|weapon|military)\b'),
    (r'\bgraphical user interf',  "03_Programming",            None),
    (r'\buser interface design\b',"03_Programming",            None),
    (r'\bsoftware process\b',     "03_Programming",            None),
    (r'\bcomputer science\b',     "03_Programming",            r'\b(defence|weapon|military|ballistic)\b'),
    (r'\bsoftware engineering\b', "03_Programming",            r'\b(defence|weapon|military)\b'),
    (r'\bweb development\b',      "03_Programming",            None),
    (r'\bdata structures?\b',     "03_Programming",            None),
    (r'\bcloud application\b',    "03_Programming",            None),
    (r'\bauthentication\b',       "03_Programming",            None),
    (r'\bcybersecurity\b',        "03_Programming",            None),
]

def _correct_domain(title: str, subject: str, raw_domain: str) -> str:
    if raw_domain not in _VALID_DOMAINS:
        raw_domain = "00_Uncategorised"
    check = f"{title} {subject}".lower()
    for pattern, correct, avoid_if in _DOMAIN_CORRECTIONS:
        if re.search(pattern, check):
            if avoid_if and re.search(avoid_if, check):
                continue
            return correct
    return raw_domain

# ── FIX 8: canonical subject mapping ────────────────────────────────────────────
_DOMAIN_SUBJECT = {
    "01_Math":"Mathematics", "02_Physics":"Physics", "03_Programming":"Computer Science",
    "04_Electronics":"Electronics Engineering", "05_Robotics":"Robotics Engineering",
    "06_AI_ML":"Artificial Intelligence", "07_Control":"Control Systems",
    "08_Design":"Engineering Design", "09_Defence":"Defence Studies", "10_HSSM":"Humanities",
    "11_Projects":"Projects & Research", "13_MechanicalEngineering":"Mechanical Engineering",
}
_CHEM_KW = ('chemistry','chemical','acid','organic','inorganic')
_BIO_KW  = ('biology','biological','cell','physiology','genetic','ecology','plant','organism','evolution','reproduct')

def _canonical_subject(domain: str, title: str) -> str:
    if domain == "12_NaturalSciences":
        low = title.lower()
        if any(k in low for k in _CHEM_KW): return "Chemistry"
        if any(k in low for k in _BIO_KW):  return "Biology"
        return "Natural Sciences"
    return _DOMAIN_SUBJECT.get(domain, "General Studies")

# ── FIX 10: software/systems engine override ────────────────────────────────────
_SOFTWARE_TITLE_KEYWORDS = ('ros)','robot operating system','operating system','microcontroller','interfacing','vision system','programming')
def _is_software_systems_course(title: str) -> bool:
    return any(kw in title.lower() for kw in _SOFTWARE_TITLE_KEYWORDS)

_MATH_MARKERS = (
    'derivation','derive','hyperplane','matrix','matrices','eigen','gradient','equation','theorem',
    'transform','jacobian','differential','integral','optimization','transfer function','probability',
    'distribution','svd','singular value','statistic','convolution','backpropagation','regression',
    'principal component','loss function','stability criterion','bayesian','markov','dynamics','kalman',
    'kinematics','kinematic','set theory','membership','stochastic','variance','vector calculus','tensor',
)
_MARKER_PATTERNS = [re.compile(r'\b'+re.escape(m)+r'\b') for m in _MATH_MARKERS]

def _resolve_engine(domain: str, title: str, core_subtopics: list[str]) -> str:
    if _is_software_systems_course(title):
        return "remotion"
    if domain in MANIM_DOMAINS:
        return "manim"
    if domain == "10_HSSM":
        return "remotion"
    text = " ".join(core_subtopics).lower()
    hits = sum(1 for pat in _MARKER_PATTERNS if pat.search(text))
    return "manim" if hits >= 2 else "remotion"

# ── FIX 9: anchor-prefixed web_queries — deterministic, overrides LLM output ───
_GENERIC_LEAD_WORDS = {'introduction','fundamentals','basics','principles','concepts'}
# BUG FIX (found via adversarial testing): titles like "Introduction to
# Fundamentals of Basics of Advanced Robotics..." stripped the generic NOUNS
# but kept connecting words "to"/"of", which then got selected as the anchor
# since they were first in the filtered list -> anchor="to of of" (meaningless,
# reintroduces the exact ambiguous-query problem this function exists to fix).
_FILLER_WORDS = {'to','of','in','for','with','on','at','by','the','a','an','and','or'}

def _strip_nested_parens(text: str) -> str:
    """
    BUG FIX: the original single-pass re.sub(r'\\([^)]*\\)') mishandled nested
    parens — "Topic (Outer (Inner) Stuff)" left a dangling unmatched ')'.
    [^()]* (excluding both bracket chars) + looping until stable correctly
    peels nested groups from the innermost outward.
    """
    prev = text
    for _ in range(4):
        nxt = re.sub(r'\([^()]*\)', '', prev)
        if nxt == prev: break
        prev = nxt
    return prev

def _extract_anchor(title: str) -> str:
    paren = re.search(r'\(([A-Z]{2,6})\)', title)
    if paren: return paren.group(1)
    clean = _strip_nested_parens(title).strip()
    clean = re.sub(r'\s{2,}', ' ', clean)   # collapse double-spaces left by paren removal
    words = clean.split()
    if len(words) <= 5: return re.sub(r'[,:]', '', clean).strip()
    sig = [w for w in words
           if w.lower().rstrip('.,:') not in _GENERIC_LEAD_WORDS
           and w.lower().rstrip('.,:') not in _FILLER_WORDS]
    sig = sig if sig else words
    return re.sub(r'[,:]', '', " ".join(sig[:3])).strip()

def _extract_phrase(subtopic: str) -> str:
    s = subtopic
    if ':' in s: s = s.split(':', 1)[1]
    paren = re.findall(r'\(([^)]+)\)', s)
    base  = re.sub(r'\([^)]*\)', '', s)
    clauses = [c.strip(' .') for c in base.split(',') if c.strip()] or [base.strip()]
    def score(c):
        sc = 0
        if re.search(r'\b[A-Z]{2,}\b', c): sc += 2
        if re.search(r'[A-Z][a-z]+-[A-Z][a-z]+', c): sc += 3
        if re.search(r'\d', c): sc += 1
        sc += min(len(c)//10, 3)
        return sc
    best = max(clauses, key=score)
    if paren:
        fi = paren[0].split(',')[0].strip()
        if fi.lower() not in best.lower(): best = f"{best} ({fi})"
    return re.sub(r'\s{2,}', ' ', best).strip(' .')[:60]

def _build_web_queries(title: str, core_subtopics: list[str], subject: str) -> list[str]:
    anchor = _extract_anchor(title)
    if not core_subtopics:
        return [f"{anchor} lecture notes", f"{anchor} tutorial"]
    p0 = _extract_phrase(core_subtopics[0])
    p1 = _extract_phrase(core_subtopics[-1]) if len(core_subtopics) > 1 else p0
    q1 = f"NPTEL {anchor} {p0}"
    q2 = (f"MIT OCW {anchor} {p1}" if p1.lower() != p0.lower()
          else f"{subject} {anchor} {p1} advanced tutorial")
    return [q1, q2]

# ── FIX 11: ceiling-based chapter formula ──────────────────────────────────────
def _suggested_chapters(core_subtopics: list[str]) -> int:
    return max(1, math.ceil(len(core_subtopics) / 2))

# ── FIX 2/3: semantic-aware ID generation ──────────────────────────────────────
def normalize_title(title: str) -> str:
    t = title.lower().strip()
    t = re.sub(r'\b(and|or|of|in|for|with|&|the|a|an)\b', '', t)
    return re.sub(r'[^a-z0-9]', '', t)

def _make_id(domain: str, grade_or_sem: str, title: str) -> str:
    sem_tok = grade_or_sem.replace(" ", "").lower() if grade_or_sem else "nosem"
    slug    = safe_slug(f"{domain}_{sem_tok}_{title}")[:45].rstrip("_")
    norm    = normalize_title(title)
    suffix  = hashlib.md5(f"{domain}|{sem_tok}|{norm}".encode()).hexdigest()[:8]
    return f"{slug}_{suffix}" if slug else f"mod_{suffix}"

def _merge_sources(*sources: str) -> str:
    parts: set[str] = set()
    for s in sources:
        for part in str(s or "").split(" | "):
            p = part.strip()
            if p: parts.add(p)
    return " | ".join(sorted(parts))

# ── Schema ────────────────────────────────────────────────────────────────────
_SCHEMA = {
    "type": "ARRAY",
    "items": {
        "type": "OBJECT",
        "properties": {
            "title":              {"type": "STRING"},
            "subject":            {"type": "STRING"},
            "subject_code":       {"type": "STRING"},
            "domain":             {"type": "STRING", "enum": _VALID_DOMAINS},
            "grade_or_sem":       {"type": "STRING"},
            "tier":               {"type": "STRING",  "enum": ["Junior","Senior","Engineering"]},
            "complexity":         {"type": "STRING",  "enum": ["Basic","Intermediate","Advanced","Mixed"]},
            "priority":           {"type": "INTEGER"},
            "engine":             {"type": "STRING",  "enum": ["manim","remotion"]},
            "core_subtopics":     {"type": "ARRAY",   "items": {"type": "STRING"},
                                   "description": "3-7 granular sub-topics from syllabus text for synthesis engine"},
            "suggested_chapters": {"type": "INTEGER",
                                   "description": "1-4: how many video chapters to split this module into"},
            "web_queries":        {"type": "ARRAY",   "items": {"type": "STRING"}},
        },
        "required": [
            "title","subject","subject_code","domain","grade_or_sem",
            "tier","complexity","priority","engine",
            "core_subtopics","suggested_chapters","web_queries",
        ],
    },
}

_SYSTEM = """You are an academic syllabus parser for an Indian B.Tech Robotics & AI university programme.

EXTRACT only real academic modules, subjects, courses, or teaching units.

SKIP entirely — do NOT create entries for:
- University/institution names, addresses ("Ramdeobaba University", "Nagpur-440013")
- Programme/document titles ("Programme Scheme & Syllabi", "Scheme of Teaching")
- Administrative headers ("Course Objectives", "Course Outcomes", "Teaching Scheme")
- Track/choice labels ("OR Internship Track", "Open Elective", "Standard Track",
  "Specialization Schemes", "Choose either Track A or B")
- Textbook/reference lists, department/school names, column headers, evaluation text
- Degree strings ("B.Tech","B. Tech."), practical instructions, programme structure notes

CRITICAL: Extract ONLY text a teacher delivers a lecture on. Administrative text -> SKIP.

DOMAIN ASSIGNMENT:
01_Math: Mathematics, Calculus, Algebra, Statistics, Probability, Trigonometry
02_Physics: Physics, Applied Physics, Mechanics (classical), Optics, Electromagnetism
03_Programming: CS, Programming, Algorithms, DSA, Software, Web Dev, Databases,
                Networking, Cybersecurity, HCI, GUI, Cloud Computing, OS, Compiler
04_Electronics: Electronics, Electrical, Circuits, Microcontrollers, Embedded, IoT, FPGA
05_Robotics: Robotics, Path Planning, ROS, SLAM, Manipulators, Autonomous Systems
06_AI_ML: AI, Machine Learning, Deep Learning, Neural Networks, NLP, Computer Vision, RL
07_Control: Control Systems, PID, Transfer Functions, Bode Plot, State Space, Stability
08_Design: Engineering Design, CAD, 3D Printing, Fabrication, Engineering Graphics
09_Defence: Defence Technology, Military Systems, Weapon Systems, MDM (ONLY explicit defence)
10_HSSM: English, Communication, History, Geography, Social Science, Economics, Ethics
11_Projects: Project Work, Capstone, Internship, Research Methodology, Dissertation
12_NaturalSciences: Biology, Chemistry, Biochemistry, Botany, Zoology, Ecology, Genetics
13_MechanicalEngineering: Theory of Mechanisms, Elasticity, Manufacturing, Fluid Mechanics,
                           Heat Transfer, Strength of Materials, Machine Design, Turbomachinery
00_Uncategorised: ONLY if none above fits

ENGINE (evaluate per core_subtopics, not just domain):
- manim: ONLY when subtopics are dominated by mathematical proofs, vector calculus,
         coordinate transformations, differential equations, control loop graphs.
- remotion: code implementation, architectural diagrams, algorithms, UI, biological structures,
            chemical processes, asset-heavy content. Default when uncertain.
(NOTE: your engine choice and web_queries will be validated/corrected by deterministic
post-processing — focus your effort on accurate domain/tier/core_subtopics extraction.)

GRADE/TIER: "Sem 1"-"Sem 8" for B.Tech; "Class 6"-"Class 12" for school; empty if unknown
PRIORITY: 5=core exam | 4=important | 3=standard | 2=supplementary | 1=optional
CORE_SUBTOPICS: 3-7 granular elements directly from syllabus text (not generic labels)
SUGGESTED_CHAPTERS: 1-4 video chapters; most modules = 1 or 2
WEB_QUERIES: 2 specific strings for NPTEL/MIT OCW lecture notes"""


# ── Text extractors ────────────────────────────────────────────────────────────
def _extract_docx(path: Path) -> str:
    try:
        import docx
    except ImportError:
        raise ImportError("pip install python-docx")
    # CONSISTENCY FIX (found via testing): docx.Document() on a missing or
    # corrupted file raised an uncaught PackageNotFoundError. map()'s outer
    # try/except already catches this in the real pipeline (files come from
    # glob(), so they exist), but this function should be self-contained like
    # _extract_pdf already is, rather than relying on the caller's safety net.
    try:
        doc = docx.Document(str(path))
    except Exception as e:
        log.warning(f"  [yellow]Cannot open {path.name} as DOCX: {e}[/yellow]")
        return ""
    parts = []
    for para in doc.paragraphs:
        t = para.text.strip()
        if t: parts.append(t)
    for tbl in doc.tables:
        for row in tbl.rows:
            cells = []
            for cell in row.cells:
                ct = re.sub(r'\s{2,}', ' ', cell.text.strip().replace('\t', ' '))
                if ct: cells.append(ct)
            if cells: parts.append(" | ".join(cells))
    return "\n".join(parts)

def _extract_pdf(path: Path) -> str:
    try:
        import pypdf
        parts = []
        with open(path, "rb") as f:
            r = pypdf.PdfReader(f)
            for i in range(min(len(r.pages), 120)):
                try:
                    t = r.pages[i].extract_text()
                    if t and t.strip(): parts.append(t.strip())
                except Exception: continue
        text = "\n".join(parts)
        if len(text.strip()) > 200: return text
    except Exception as e:
        log.warning(f"  [yellow]pypdf failed {path.name}: {e}[/yellow]")
    try:
        from pdfminer.high_level import extract_text
        return extract_text(str(path)) or ""
    except Exception as e:
        log.warning(f"  [yellow]pdfminer failed {path.name}: {e}[/yellow]")
    return ""

def _extract_text(path: Path) -> str:
    ext = path.suffix.lower()
    if ext == ".docx": return _extract_docx(path)
    if ext == ".pdf":  return _extract_pdf(path)
    return ""

def _truncate_at_line(text: str, max_chars: int) -> str:
    if len(text) <= max_chars: return text
    cut = text.rfind('\n', 0, max_chars)
    return text[:cut] if cut > max_chars // 2 else text[:max_chars]

def _file_hash(text: str) -> str:
    return hashlib.md5(text.encode("utf-8")).hexdigest()

def _load_file_state(state_path: Path) -> dict:
    if not state_path.exists(): return {}
    try:
        s = json.loads(state_path.read_text(encoding="utf-8"))
        return s if isinstance(s, dict) else {}
    except Exception: return {}

def _save_file_state(state_path: Path, state: dict) -> None:
    tmp = state_path.with_suffix(".tmp")
    try:
        state_path.parent.mkdir(parents=True, exist_ok=True)
        tmp.write_text(json.dumps(state, indent=2), encoding="utf-8")
        tmp.replace(state_path)
    except OSError as e:
        log.warning(f"[yellow]File-state save failed: {e}[/yellow]")
    finally:
        tmp.unlink(missing_ok=True)


# ── Gemini call: ONE per file, shrink-on-retry, never re-prompt for formatting ──
def _gemini_map(raw_text: str, filename: str, client: google_genai.Client) -> list[dict]:
    if not raw_text.strip():
        return []

    config = gt.GenerateContentConfig(
        system_instruction=_SYSTEM,
        response_mime_type="application/json",
        response_schema=_SCHEMA,
        temperature=0.1,
        max_output_tokens=65536,
        thinking_config=gt.ThinkingConfig(thinking_budget=0),
        tool_config=gt.ToolConfig(
            function_calling_config=gt.FunctionCallingConfig(mode="NONE")
        ),
    )

    max_chars = 120_000
    for attempt in range(MAX_RETRIES):
        limit     = int(max_chars * (0.7 ** attempt))   # shrink 30% per retry — never resend identical oversized prompt
        text_body = _truncate_at_line(raw_text, limit)

        prompt = (
            f"Parse this syllabus and extract all academic modules.\n\n"
            f"FILENAME: {filename}\n\n"
            f"DOCUMENT TEXT:\n{'─'*60}\n"
            f"{text_body}\n"
            f"{'─'*60}\n\n"
            f"Return JSON array. Apply CRITICAL PARSING rules exactly."
        )

        try:
            response = client.models.generate_content(
                model=GEMINI_FLASH_MODEL, contents=prompt, config=config,
            )
            raw_out = safe_gemini_text(response)
            if not raw_out:
                log.warning(f"  [yellow]Empty response attempt {attempt+1}[/yellow]")
                time.sleep(calc_backoff(attempt)); continue

            parsed = robust_json_parse(raw_out)
            if parsed is None:
                log.warning(f"  [yellow]JSON parse failed attempt {attempt+1} (text={len(text_body):,} chars)[/yellow]")
                time.sleep(calc_backoff(attempt)); continue

            if not isinstance(parsed, list):
                for k in ("modules","items","data","results"):
                    if isinstance(parsed, dict) and isinstance(parsed.get(k), list):
                        parsed = parsed[k]; break
            if not isinstance(parsed, list):
                log.warning(f"  [yellow]Not array: {type(parsed).__name__}[/yellow]")
                time.sleep(calc_backoff(attempt)); continue

            valid = [
                m for m in parsed
                if isinstance(m, dict)
                and len(str(m.get("title","")).strip()) >= 3
                and len(str(m.get("subject","")).strip()) >= 2
            ]
            log.info(
                f"  [green]Gemini mapped {len(valid)} modules from {filename}[/green]"
                + (f" [dim](dropped {len(parsed)-len(valid)} empty)[/dim]"
                   if len(parsed) != len(valid) else "")
            )
            return valid

        except Exception as e:
            es = str(e).lower()
            if "429" in es or "resource_exhausted" in es or "quota" in es:
                wait = min(60*(attempt+1), 300)
                log.warning(f"  [yellow]Rate limit — waiting {wait}s[/yellow]")
                time.sleep(wait); continue
            if ("500" in es or "503" in es or "unavailable" in es) and attempt < MAX_RETRIES-1:
                time.sleep(calc_backoff(attempt)); continue
            log.error(f"  [red]Gemini error {filename}: {e}[/red]")
            if attempt < MAX_RETRIES-1:
                time.sleep(calc_backoff(attempt)); continue
            return []
    return []


# ── Entry builder — ALL deterministic post-processing applied here ─────────────
def _build_entry(raw: dict, src: Path, existing: dict) -> Optional[dict]:
    title = str(raw.get("title","")).strip()
    if not title: return None

    subject_raw  = str(raw.get("subject", title)).strip() or title
    subject_code = str(raw.get("subject_code","")).strip()
    raw_domain   = str(raw.get("domain","")).strip()
    domain       = _correct_domain(title, subject_raw, raw_domain)
    if domain not in _VALID_DOMAINS: domain = "00_Uncategorised"

    grade_or_sem = str(raw.get("grade_or_sem","") or "").strip()
    tier         = str(raw.get("tier","")).strip()
    if tier not in ("Junior","Senior","Engineering"): tier = "Engineering"

    complexity = str(raw.get("complexity","Mixed")).strip()
    if complexity not in ("Basic","Intermediate","Advanced","Mixed"): complexity = "Mixed"

    try:    priority = max(1, min(5, int(raw.get("priority",3))))
    except: priority = 3

    cst            = raw.get("core_subtopics",[])
    core_subtopics = [str(s).strip() for s in cst if str(s).strip()][:10] if isinstance(cst,list) else []

    # FIX 10: engine ALWAYS resolved deterministically — Gemini's raw value discarded
    engine = _resolve_engine(domain, title, core_subtopics)

    # FIX 11: ceiling chapter formula — Gemini's raw value discarded
    suggested_chapters = _suggested_chapters(core_subtopics)

    # FIX 8: canonical subject; original descriptive label preserved as track
    canonical_subj = _canonical_subject(domain, title)
    track = subject_raw if subject_raw != title else (
        "B.Tech Robotics & AI" if tier == "Engineering" else canonical_subj
    )

    # FIX 9: anchor-prefixed web_queries — Gemini's raw value discarded
    web_queries = _build_web_queries(title, core_subtopics, canonical_subj)

    entry_id = _make_id(domain, grade_or_sem, title)

    source_files: list[str] = []
    sp = STAGING_DIR / tier / safe_slug(subject_raw)
    if sp.exists():
        try:    source_files = [str(f) for f in sp.iterdir() if f.is_file()]
        except OSError: pass

    entry = {
        "id": entry_id, "title": title, "subject": canonical_subj,
        "subject_code": subject_code, "domain": domain, "tier": tier,
        "grade_or_sem": grade_or_sem, "complexity": complexity, "priority": priority,
        "engine": engine, "core_subtopics": core_subtopics,
        "suggested_chapters": suggested_chapters, "web_queries": web_queries,
        "source_files": source_files, "source_docx": str(src),
        "content": "", "references": [], "pyqs": [], "track": track,
        "prerequisite_context": [],   # filled by _compute_prerequisites() after full parse
    }

    matched = existing.get(entry_id)
    if matched is None:
        norm = normalize_title(title)
        for eid, p in existing.items():
            if (normalize_title(p.get("title","")) == norm
                    and p.get("domain","") == domain
                    and str(p.get("grade_or_sem","") or "") == grade_or_sem):
                matched = p; break
    if matched:
        entry["content"]    = matched.get("content","")
        entry["references"] = matched.get("references",[]) if isinstance(matched.get("references"),list) else []
        pyqs_val             = matched.get("pyqs",[])
        entry["pyqs"]        = pyqs_val if isinstance(pyqs_val,list) else []

    return entry


# ── FIX 12: prerequisite_context — final cross-entry pass, zero API cost ───────
_TIER_RANK = {"Junior": 0, "Senior": 1, "Engineering": 2}

def _grade_rank(tier: str, grade_or_sem: str) -> int:
    gs = grade_or_sem or ""
    if tier == "Junior":
        if "6-8" in gs: return 0
        if gs == "Class 9" or "9-10" in gs: return 1
        if gs == "Class 10": return 2
        return 1
    if tier == "Senior": return 0
    if tier == "Engineering":
        nums = re.findall(r'\d+', gs)
        return int(nums[0]) if nums else 0
    return 0

def _compute_prerequisites(final: list[dict]) -> None:
    by_subject: dict[str, list[dict]] = {}
    for m in final:
        by_subject.setdefault(m["subject"], []).append(m)
    for subject, entries in by_subject.items():
        ranked = sorted(entries, key=lambda m: (_TIER_RANK[m["tier"]], _grade_rank(m["tier"], m["grade_or_sem"])))
        for i, m in enumerate(ranked):
            rank = (_TIER_RANK[m["tier"]], _grade_rank(m["tier"], m["grade_or_sem"]))
            earlier = [e for e in ranked[:i] if (_TIER_RANK[e["tier"]], _grade_rank(e["tier"], e["grade_or_sem"])) < rank]
            m["prerequisite_context"] = [
                f"{e['title']} ({e['tier']}, {e['grade_or_sem']})" for e in earlier[-4:]
            ]


# ── CurriculumMapper ───────────────────────────────────────────────────────────
class CurriculumMapper:
    def __init__(self, gemini_client: google_genai.Client) -> None:
        self.client = gemini_client

    def map(self, curriculum_dir: Path, output: Path, clean: bool = False, force: bool = False) -> list[dict]:
        curriculum_dir.mkdir(parents=True, exist_ok=True)
        docx_files = sorted(curriculum_dir.glob("*.docx"))
        pdf_files  = [p for p in sorted(curriculum_dir.glob("*.pdf"))
                      if not (curriculum_dir/(p.stem+".docx")).exists()]
        all_files = docx_files + pdf_files

        existing: dict[str,dict] = {}
        if not clean and output.exists():
            try:
                prev = json.loads(output.read_text(encoding="utf-8"))
                if isinstance(prev,list):
                    existing = {e["id"]:e for e in prev if isinstance(e,dict) and e.get("id")}
                    log.info(f"[cyan]{len(existing)} existing entries preserved[/cyan]")
            except Exception as e:
                log.warning(f"[yellow]Cannot load manifest: {e} — fresh start[/yellow]")
        elif clean:
            log.info("[yellow]--clean: ignoring existing manifest[/yellow]")

        if not all_files:
            log.error(f"[red]No DOCX/PDF in {curriculum_dir}[/red]"); return []

        # API-BLEED FIX: skip Gemini call entirely for files whose content hash
        # matches the last successful run AND whose entries already exist in
        # the manifest. Without this, re-running the mapper after adding ONE
        # new file re-parses every file via Gemini again — wasted tokens.
        state_path  = curriculum_dir.parent / "curriculum_mapper_state.json"
        file_state  = {} if (clean or force) else _load_file_state(state_path)
        new_file_state: dict[str, str] = {}

        log.info(f"\n[bold]Phase 0[/bold] — {len(docx_files)} DOCX + {len(pdf_files)} PDF")

        all_modules: list[dict] = []
        gemini_calls = 0
        skipped_unchanged = 0
        for f in all_files:
            log.info(f"\n  [cyan]Processing:[/cyan] {f.name}")
            try:
                raw_text = _extract_text(f)
            except Exception as e:
                log.error(f"  [red]Extraction failed {f.name}: {e}[/red]"); continue
            if not raw_text or len(raw_text.strip()) < 50:
                log.warning(f"  [yellow]Too little text — skipping {f.name}[/yellow]"); continue

            file_key  = str(f.resolve())
            this_hash = _file_hash(raw_text)
            prior_entries = [e for e in existing.values() if e.get("source_docx") == str(f)]

            if (not force and not clean and file_state.get(file_key) == this_hash
                    and prior_entries):
                log.info(f"  [dim]Unchanged since last run — reusing {len(prior_entries)} "
                         f"entries, no Gemini call[/dim]")
                all_modules.extend(prior_entries)
                new_file_state[file_key] = this_hash
                skipped_unchanged += 1
                continue

            log.info(f"  [dim]Extracted {len(raw_text):,} chars[/dim]")
            modules = _gemini_map(raw_text, f.name, self.client)
            gemini_calls += 1
            if not modules:
                log.warning(f"  [yellow]No modules from Gemini for {f.name}[/yellow]"); continue

            for raw in modules:
                try:
                    entry = _build_entry(raw, f, existing)
                    if entry: all_modules.append(entry)
                except Exception as e:
                    log.error(f"  [red]Entry build failed '{raw.get('title','?')}': {e}[/red]")
            new_file_state[file_key] = this_hash

        fresh_ids = {m["id"] for m in all_modules}
        carried   = [e for eid,e in existing.items() if eid not in fresh_ids]
        if carried: log.info(f"[dim]Carrying forward {len(carried)} enriched entries[/dim]")
        all_modules.extend(carried)

        # Pass 1: ID-based dedup
        deduped: dict[str,dict] = {}
        for m in all_modules:
            mid = m.get("id")
            if not mid: continue
            if mid not in deduped: deduped[mid] = m
            else: deduped[mid]["source_docx"] = _merge_sources(deduped[mid].get("source_docx",""), m.get("source_docx",""))

        # Pass 2: semantic dedup (catches & vs and slug divergence)
        seen_sem: dict[str,int] = {}
        final: list[dict] = []
        for m in deduped.values():
            sk = m.get("domain","") + "|" + str(m.get("grade_or_sem","") or "") + "|" + normalize_title(m.get("title",""))
            if sk not in seen_sem:
                seen_sem[sk] = len(final); final.append(m)
            else:
                kept = final[seen_sem[sk]]
                kept["source_docx"] = _merge_sources(kept.get("source_docx",""), m.get("source_docx",""))
                for field in ("content",):
                    if not kept.get(field) and m.get(field): kept[field] = m[field]
                for field in ("references","pyqs"):
                    kv, mv = kept.get(field) or [], m.get(field) or []
                    if isinstance(mv,list) and len(mv) > len(kv): kept[field] = mv
                if len(m.get("core_subtopics",[])) > len(kept.get("core_subtopics",[])):
                    kept["core_subtopics"] = m["core_subtopics"]
                if not kept.get("subject_code") and m.get("subject_code"): kept["subject_code"] = m["subject_code"]

        if not final:
            log.error("[red]No modules produced.[/red]"); return []

        def _sk(m: dict) -> tuple:
            gs = str(m.get("grade_or_sem") or "")
            n  = int(re.findall(r'\d+', gs)[0]) if re.findall(r'\d+', gs) else 0
            return (-m.get("priority",3), 13+n if "Sem" in gs else n, m.get("title","").lower())
        final.sort(key=_sk)

        # FIX 12: prerequisite_context — whole-manifest pass, zero API calls
        _compute_prerequisites(final)

        tmp = output.with_suffix(".tmp")
        try:
            output.parent.mkdir(parents=True, exist_ok=True)
            tmp.write_text(json.dumps(final,indent=2,ensure_ascii=False), encoding="utf-8")
            tmp.replace(output)
            _save_file_state(state_path, new_file_state)
        except OSError as e:
            tmp.unlink(missing_ok=True); log.error(f"[red]Write failed: {e}[/red]"); raise

        tier_ct: dict[str,int] = {}
        for m in final:
            t = m.get("tier","?"); tier_ct[t] = tier_ct.get(t,0)+1

        from rich.panel import Panel
        console.print(Panel(
            f"[bold]Total modules  : {len(final)}[/bold]\n\nBy tier:\n"
            + "\n".join(f"  {t:18}: {n}" for t,n in sorted(tier_ct.items()))
            + f"\n\nGemini calls used: {gemini_calls}  |  Files unchanged (skipped, $0): {skipped_unchanged}"
            + f"\nOutput: {output}",
            title="[bold green]Phase 0 Complete[/bold green]",
            border_style="green",
        ))
        return final


def main() -> None:
    from pve.safety import require_legacy_paid_execution
    require_legacy_paid_execution()
    p = argparse.ArgumentParser(prog="curriculum_mapper.py")
    p.add_argument("--curriculum-dir", type=Path, default=CURRICULUM_DIR)
    p.add_argument("--output",         type=Path, default=MASTER_MANIFEST)
    p.add_argument("--clean",          action="store_true",
                   help="Ignore existing manifest — full fresh parse, clears zombies")
    p.add_argument("--force",          action="store_true",
                   help="Re-call Gemini for all files even if unchanged since last run (costs extra tokens)")
    args   = p.parse_args()
    client = google_genai.Client(api_key=get_env("GEMINI_API_KEY"))
    CurriculumMapper(gemini_client=client).map(args.curriculum_dir, args.output, clean=args.clean, force=args.force)

if __name__ == "__main__":
    main()
